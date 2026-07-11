#!/usr/bin/env python3
"""
Inspector de archivos Excel del proceso Benchmark Mensual.

Dado uno o más archivos .xlsb / .xlsm / .xlsx, produce artefactos de análisis
(sin exponer los binarios pesados):

  - <salida>/vba/<archivo>__<modulo>.vba   -> código VBA de cada módulo
  - <salida>/sheets/<archivo>.md           -> inventario de hojas: dimensiones,
                                              cabeceras y primeras filas de muestra
  - <salida>/index.md                      -> resumen general

Uso:
    python3 tools/inspect_excel.py <ruta_a_archivos_o_carpeta> --out analisis \
        [--sample-rows 30]

Notas:
  - .xlsb: valores legibles con pyxlsb (no evalúa fórmulas). VBA con oletools.
  - .xlsm/.xlsx: estructura/fórmulas con openpyxl. VBA con oletools.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path


def extract_vba(path: Path, out_dir: Path) -> list[str]:
    """Extrae módulos VBA a archivos de texto. Devuelve lista de módulos."""
    from oletools.olevba import VBA_Parser

    modules: list[str] = []
    try:
        vp = VBA_Parser(str(path))
    except Exception as e:  # noqa: BLE001
        return [f"(no se pudo abrir para VBA: {e})"]
    if not vp.detect_vba_macros():
        vp.close()
        return ["(sin macros VBA)"]

    vba_dir = out_dir / "vba"
    vba_dir.mkdir(parents=True, exist_ok=True)
    for _fn, _stream, vba_name, vba_code in vp.extract_macros():
        safe = "".join(c if c.isalnum() or c in "._-" else "_" for c in vba_name)
        dest = vba_dir / f"{path.stem}__{safe}.vba"
        dest.write_text(vba_code or "", encoding="utf-8", errors="replace")
        loc = (vba_code or "").count("\n") + 1
        modules.append(f"{vba_name} ({loc} líneas) -> {dest.name}")
    vp.close()
    return modules


def _fmt_cell(v) -> str:
    if v is None:
        return ""
    s = str(v).replace("\n", " ").replace("|", "/").strip()
    return s[:40]


def inspect_xlsb(path: Path, sample_rows: int) -> list[str]:
    from pyxlsb import open_workbook

    lines: list[str] = []
    with open_workbook(str(path)) as wb:
        lines.append(f"Hojas ({len(wb.sheets)}): {', '.join(wb.sheets)}\n")
        for name in wb.sheets:
            lines.append(f"\n### Hoja: `{name}`")
            with wb.get_sheet(name) as sheet:
                rows_out, ncols, nrows = [], 0, 0
                for r_idx, row in enumerate(sheet.rows()):
                    nrows = r_idx + 1
                    if r_idx <= sample_rows:
                        vals = [_fmt_cell(c.v) for c in row]
                        ncols = max(ncols, len(vals))
                        rows_out.append(vals)
                    if r_idx > sample_rows and r_idx % 50000 == 0:
                        pass  # sigue contando filas
                lines.append(f"- Filas (aprox, muestra): {nrows} · Columnas: {ncols}")
                if rows_out:
                    lines.append("\n```")
                    for vals in rows_out[: sample_rows + 1]:
                        lines.append(" | ".join(vals))
                    lines.append("```")
    return lines


def inspect_openpyxl(path: Path, sample_rows: int) -> list[str]:
    import openpyxl

    lines: list[str] = []
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=False)
    lines.append(f"Hojas ({len(wb.sheetnames)}): {', '.join(wb.sheetnames)}\n")
    for name in wb.sheetnames:
        ws = wb[name]
        lines.append(f"\n### Hoja: `{name}`")
        lines.append(f"- Dimensiones: {ws.max_row} filas x {ws.max_column} columnas")
        rows_out = []
        for r_idx, row in enumerate(ws.iter_rows(max_row=sample_rows + 1, values_only=True)):
            rows_out.append([_fmt_cell(v) for v in row])
        if rows_out:
            lines.append("\n```")
            for vals in rows_out:
                lines.append(" | ".join(vals))
            lines.append("```")
    wb.close()
    return lines


def process(path: Path, out_dir: Path, sample_rows: int) -> str:
    ext = path.suffix.lower()
    header = [f"# {path.name}", f"- Tamaño: {path.stat().st_size/1e6:.1f} MB", f"- Tipo: {ext}"]

    # VBA
    header.append("\n## Módulos VBA")
    for m in extract_vba(path, out_dir):
        header.append(f"- {m}")

    # Hojas
    header.append("\n## Inventario de hojas")
    try:
        if ext == ".xlsb":
            header += inspect_xlsb(path, sample_rows)
        elif ext in (".xlsm", ".xlsx"):
            header += inspect_openpyxl(path, sample_rows)
        else:
            header.append(f"(extensión no soportada: {ext})")
    except Exception as e:  # noqa: BLE001
        header.append(f"(error leyendo hojas: {e})")

    sheets_dir = out_dir / "sheets"
    sheets_dir.mkdir(parents=True, exist_ok=True)
    dest = sheets_dir / f"{path.stem}.md"
    dest.write_text("\n".join(header), encoding="utf-8")
    return dest.name


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("target", help="Archivo o carpeta con .xlsb/.xlsm/.xlsx")
    ap.add_argument("--out", default="analisis", help="Carpeta de salida")
    ap.add_argument("--sample-rows", type=int, default=30)
    args = ap.parse_args()

    target = Path(args.target)
    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)

    if target.is_dir():
        files = sorted(
            p for p in target.iterdir()
            if p.suffix.lower() in (".xlsb", ".xlsm", ".xlsx") and not p.name.startswith("~$")
        )
    else:
        files = [target]

    if not files:
        print(f"No se encontraron archivos Excel en {target}", file=sys.stderr)
        return 1

    index = ["# Índice de análisis de archivos Excel\n"]
    for f in files:
        print(f"Procesando {f.name} ...", file=sys.stderr)
        name = process(f, out_dir, args.sample_rows)
        index.append(f"- [{f.name}](sheets/{name})")
    (out_dir / "index.md").write_text("\n".join(index), encoding="utf-8")
    print(f"Listo. Resultados en {out_dir}/", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
